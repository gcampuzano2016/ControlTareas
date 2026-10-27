# -*- coding: utf-8 -*-
"""
Genera el script de migracion del modulo de Bonos desde el Excel de septiembre.

POR QUE EXISTE ESTE ARCHIVO Y NO UN .sql CON LOS DATOS DENTRO:
el Excel trae nombres, codigos SAP, cedulas y montos por persona, y este
repositorio es PUBLICO. Se versiona la logica de la migracion -este archivo,
que se puede revisar- y no su resultado, que son datos personales.

La salida va a docs/sql/carga-generada/, que esta en .gitignore.

Uso:
    python -I docs/sql/generar-migracion-bonos.py

El Excel se espera en RevisionBonos/9. Septiembre.xlsx, que tambien esta
ignorado. Spec: docs/superpowers/specs/2026-10-20-bonos-design.md, seccion 7.

Que migra:
  1. Casi nada del catalogo: Talento Humano lo armo a mano en produccion desde
     el 2026-10-23, con otros nombres. Cada tipo del Excel va al tipo de
     produccion de DESTINO; solo se crean los de NUEVOS y se marcan por semana
     los de SEMANALES. Nombres, valores y conceptos del catalogo no se tocan.
  2. Septiembre 2026 como periodo Cerrado, con un bono Aprobado de
     Origen = Migracion por cada fila del Excel (salvo la de $0). Su concepto
     es el que se pago en septiembre (CONCEPTO), no el del catalogo.
     Si falta un tipo de destino, si no coincide en "por semana" o si una
     semana ya la paga otro bono de la misma persona y tipo, no migra nada.
  3. Abril 2025 - agosto 2026 como periodos Historico con su total.

Que NO hace el generador y si el script: encontrar a la persona. El Excel
solo trae empresa e identificacion (DOS: Cod_Sap; AGILITY: cedula); el
script la busca entre las personas activas de Sp_RTA_BonoPersonasActivas y,
si alguna fila no aparece (o aparece dos veces, o con otra empresa), no
migra nada y lista las filas por numero.

Lo que el generador imprime -ajustes de semanas, la fila omitida- lleva
numero de fila y empresa, nunca nombre, cedula ni codigo SAP.
"""
import datetime
import difflib
import os
import re
import sys
import unicodedata
from collections import Counter
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
EXCEL = os.path.join(RAIZ, 'RevisionBonos', '9. Septiembre.xlsx')
SALIDA_DIR = os.path.join(RAIZ, 'docs', 'sql', 'carga-generada')
SALIDA = os.path.join(SALIDA_DIR, 'bonos-migracion.sql')

HOJA_SEPTIEMBRE = 'Sep_2026'
FILA_DESDE, FILA_HASTA = 5, 47
HOJA_HISTORICO = 'vista_mensual'

ANIO, MES = 2026, 9
TOTAL_SEPTIEMBRE = Decimal('3195.00')
# Primer mes y ultimo mes historicos, inclusive: 17 meses.
HISTORICO_DESDE = (2025, 4)
HISTORICO_HASTA = (2026, 8)
HISTORICOS_ESPERADOS = 17

# La primera semana marcable del modulo (spec, seccion 2 y decision 11).
PRIMER_LUNES = date(2026, 8, 3)

USUARIO = 'migracion-bonos'
MOTIVO_AJUSTE = 'Migracion desde el Excel de septiembre'

MESES = {'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6,
         'julio': 7, 'agosto': 8, 'septiembre': 9, 'octubre': 10, 'noviembre': 11,
         'diciembre': 12}

CONCEPTO = {'DOS': '1M34', 'AGILITY': '1M40'}
CONCEPTO_CUMPLIMIENTO = {'DOS': '1M30', 'AGILITY': '1M37'}

# (Nombre, Familia, ValorSugerido o None, PorSemana). Nombre con su tilde
# real: el script la escribe con NCHAR.
CATALOGO = [
    ('Standby Cisco UIO', 'Standby', Decimal('37.50'), True),
    ('Standby Cisco GYE', 'Standby', Decimal('37.50'), True),
    ('Standby F5 UIO', 'Standby', Decimal('37.50'), True),
    ('Standby F5 GYE', 'Standby', Decimal('37.50'), True),
    ('Standby Data Center UIO', 'Standby', Decimal('37.50'), True),
    ('Standby Data Center GYE', 'Standby', Decimal('37.50'), True),
    ('Standby Microsoft', 'Standby', Decimal('37.50'), True),
    ('Standby GPF', 'Standby', Decimal('40.00'), True),
    ('Standby GPF fijo', 'Standby', Decimal('50.00'), False),
    ('Standby Difare fijo', 'Standby', Decimal('50.00'), False),
    ('Contrato ESPOL', 'Contrato', Decimal('100.00'), False),
    ('Contrato GPF', 'Contrato', Decimal('60.00'), False),
    ('Contrato Belcorp', 'Contrato', Decimal('50.00'), False),
    (u'Contrato Nestlé', 'Contrato', None, False),
    ('Contrato Linde', 'Contrato', None, False),
    ('Cumplimiento', 'Otro', None, False),
    ('Nuevo cargo', 'Otro', None, False),
]
TIPOS = {t[0]: t for t in CATALOGO}

# CATALOGO solo sirve para leer el Excel (sugerido, si es por semana). Los
# tipos de verdad son los del catalogo de produccion, que Talento Humano armo a
# mano desde el 2026-10-23 con otros nombres: cada tipo del Excel va al suyo,
# para no duplicar el catalogo. Donde el nombre de produccion lleva el monto, se
# elige por el valor de la fila; un monto sin destino detiene el generador.
# Correspondencia confirmada por el usuario el 2026-10-27.
DESTINO = {
    'Standby Cisco UIO': 'Bono Standby Networking',
    'Standby Cisco GYE': 'Bono Standby Networking',
    'Standby F5 UIO': 'Bono Standby F5',
    'Standby F5 GYE': 'Bono Standby F5',
    'Standby Data Center UIO': 'Bono Standby Data Center',
    'Standby Data Center GYE': 'Bono Standby Data Center',
    'Standby Microsoft': 'Bono Standby Software',
    'Standby GPF': 'Bono Standby GPF',
    'Standby GPF fijo': 'Bono Administrativo GPF',
    'Standby Difare fijo': 'Bono de standby Difare 50',
    'Contrato ESPOL': 'Bono Contrato ESPOL 100',
    'Contrato GPF': 'Bono Contrato GPF 60',
    'Contrato Belcorp': 'Bono Contrato Belcorp 50',
    u'Contrato Nestlé': {Decimal('200.00'): u'Bono Contrato Nestlé 200',
                         Decimal('60.00'): u'Bono Contrato Nestlé 60'},
    'Contrato Linde': {Decimal('70.00'): 'Bono Contrato Linde 70',
                       Decimal('100.00'): 'Bono Contrato Linde 100'},
    'Cumplimiento': {Decimal('230.00'): 'Bono Cumplimiento 230'},
    'Nuevo cargo': 'Bono Nuevo cargo',
}

# Los unicos tipos que crea la migracion, con los conceptos de CONCEPTO:
# (Nombre, Familia, ValorSugerido o None, PorSemana).
NUEVOS = [('Bono Nuevo cargo', 'Otro', None, False)]

# Tipos de produccion que quedaron como no semanales y se pagan por semana
# (sus filas de septiembre traen semanas): la migracion los marca semanales.
SEMANALES = ['Bono Standby Networking']


class Detener(Exception):
    """Un dato que no se puede interpretar con seguridad. No se genera nada."""


def plano(texto):
    """Minusculas, sin tildes, espacios simples."""
    t = unicodedata.normalize('NFKD', str(texto or ''))
    t = ''.join(c for c in t if not unicodedata.combining(c)).lower()
    return re.sub(r'\s+', ' ', t).strip()


def semanas_de(observacion):
    """'2 semanas standby...' -> 2; sin numero de semanas -> None."""
    m = re.search(r'(\d+)\s*semanas?\b', plano(observacion))
    return int(m.group(1)) if m else None


def tipo_de(observacion):
    """Traduce una OBSERVACION del Excel a un nombre del catalogo, o None.

    'cotrato' es una errata de 'contrato'. 'bono standby gpf' con N semanas
    es el standby GPF por semana; sin semanas, el de valor fijo (decision 3:
    son conceptos distintos)."""
    t = plano(observacion).replace('cotrato', 'contrato')
    con_semanas = semanas_de(observacion) is not None
    if 'standby' in t:
        if 'cisco' in t:
            return 'Standby Cisco UIO' if 'uio' in t else ('Standby Cisco GYE' if 'gye' in t else None)
        if 'f5' in t:
            return 'Standby F5 UIO' if 'uio' in t else ('Standby F5 GYE' if 'gye' in t else None)
        if 'data center' in t:
            return ('Standby Data Center UIO' if 'uio' in t
                    else ('Standby Data Center GYE' if 'gye' in t else None))
        if 'microsoft' in t:
            return 'Standby Microsoft'
        if 'gpf' in t:
            return 'Standby GPF' if con_semanas else 'Standby GPF fijo'
        if 'difare' in t:
            return None if con_semanas else 'Standby Difare fijo'
        return None
    if 'contrato' in t:
        for clave, nombre in (('nestle', u'Contrato Nestlé'), ('linde', 'Contrato Linde'),
                              ('espol', 'Contrato ESPOL'), ('gpf', 'Contrato GPF'),
                              ('belcorp', 'Contrato Belcorp')):
            if clave in t:
                return nombre
        return None
    if 'cumplimiento' in t:
        return 'Cumplimiento'
    if 'nuevo cargo' in t:
        return 'Nuevo cargo'
    return None


def destino_de(tipo, valor, fila):
    """El tipo de produccion de una fila: por nombre, o por monto si el nombre
    de produccion lo lleva."""
    d = DESTINO.get(tipo)
    if isinstance(d, dict):
        d = d.get(valor)
    if d is None:
        raise Detener('fila %d: el tipo %s con valor %s no tiene tipo de produccion asignado.'
                      % (fila, plano(tipo), valor))
    return d


def mes_de(palabra, fila, avisos):
    """Nombre de mes en espanol, tolerando erratas ('septeimbre')."""
    if palabra in MESES:
        return MESES[palabra]
    parecido = difflib.get_close_matches(palabra, list(MESES), n=1, cutoff=0.75)
    if not parecido:
        raise Detener('fila %d: el mes "%s" de FECHAS no se reconoce.' % (fila, palabra))
    avisos.append('fila %d: el mes "%s" se leyo como "%s".' % (fila, palabra, parecido[0]))
    return MESES[parecido[0]]


RANGO = re.compile(r'^del (\d{1,2})(?: de ([a-z]+))? al (\d{1,2}) de ([a-z]+)$')


def lunes_de(fechas, n, fila, avisos):
    """Los lunes que paga una fila.

    FECHAS trae una semana por linea ('Del 3 al 10 de agosto', 'Del 31 de
    agosto al 6 de septiembre'). Escritas a mano, no siempre van de lunes a
    domingo: cada linea se ajusta al lunes de su fecha de inicio (spec, 7:
    'las semanas escritas a mano se ajustan a su lunes'). Si eso no da N
    lunes distintos -N es el numero de semanas de la observacion-, se toman
    N lunes consecutivos desde el lunes de inicio y se avisa."""
    lineas = [plano(x) for x in str(fechas or '').splitlines() if plano(x)]
    if not lineas:
        raise Detener('fila %d: la observacion dice %d semana(s) pero FECHAS esta vacia.' % (fila, n))
    lunes = []
    for linea in lineas:
        m = RANGO.match(linea)
        if not m:
            raise Detener('fila %d: no se entiende la linea de FECHAS "%s".' % (fila, linea))
        d1, m1, d2, m2 = int(m.group(1)), m.group(2), int(m.group(3)), m.group(4)
        mes_fin = mes_de(m2, fila, avisos)
        if m1:
            mes_ini = mes_de(m1, fila, avisos)
        else:
            mes_ini = mes_fin if d1 <= d2 else mes_fin - 1
        try:
            inicio, fin = date(ANIO, mes_ini, d1), date(ANIO, mes_fin, d2)
        except ValueError:
            raise Detener('fila %d: la linea de FECHAS "%s" no es una fecha valida.' % (fila, linea))
        if fin < inicio:
            raise Detener('fila %d: en "%s" el fin es anterior al inicio.' % (fila, linea))
        lun = inicio - timedelta(days=inicio.weekday())
        if inicio != lun or fin != lun + timedelta(days=6):
            avisos.append('fila %d: "%s" (%s a %s) se toma como la semana del lunes %s.'
                          % (fila, linea, inicio.isoformat(), fin.isoformat(), lun.isoformat()))
        lunes.append(lun)
    distintos = sorted(set(lunes))
    if len(distintos) != n:
        consecutivos = [distintos[0] + timedelta(weeks=i) for i in range(n)]
        avisos.append('fila %d: FECHAS da %d semana(s) y la observacion dice %d; se toman %d lunes '
                      'consecutivos desde %s: %s.'
                      % (fila, len(distintos), n, n, distintos[0].isoformat(),
                         ', '.join(x.isoformat() for x in consecutivos)))
        distintos = consecutivos
    if distintos[0] < PRIMER_LUNES:
        movidos = [PRIMER_LUNES + timedelta(weeks=i) for i in range(n)]
        if len(set(movidos) | set(x for x in distintos if x >= PRIMER_LUNES)) != n:
            raise Detener('fila %d: hay semanas anteriores al %s y moverlas cambia la cantidad.'
                          % (fila, PRIMER_LUNES.isoformat()))
        avisos.append('fila %d: semanas anteriores al %s; se mueven a %s.'
                      % (fila, PRIMER_LUNES.isoformat(), ', '.join(x.isoformat() for x in movidos)))
        distintos = movidos
    return distintos


def dinero(v, fila, columna):
    if v is None or (isinstance(v, str) and not v.strip()):
        raise Detener('fila %s: la columna %s esta vacia.' % (fila, columna))
    try:
        return Decimal(str(v).strip()).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    except Exception:
        raise Detener('fila %s: la columna %s no es un numero.' % (fila, columna))


def identificacion(empresa, crudo, fila):
    """DOS: Cod_Sap tal cual. AGILITY: cedula rellenada a 10 (Excel guarda
    como numero las que empiezan con cero y les quita el cero)."""
    if crudo is None:
        raise Detener('fila %d: la columna ID SAP esta vacia.' % fila)
    if isinstance(crudo, float):
        if crudo != int(crudo):
            raise Detener('fila %d: la columna ID SAP no es un entero.' % fila)
        crudo = int(crudo)
    texto = str(crudo).strip()
    if not texto.isdigit():
        raise Detener('fila %d: la columna ID SAP no son solo digitos.' % fila)
    if empresa == 'AGILITY':
        if len(texto) > 10:
            raise Detener('fila %d: la cedula tiene mas de 10 digitos.' % fila)
        texto = texto.zfill(10)
    elif len(texto) > 32:
        raise Detener('fila %d: el codigo SAP supera 32 caracteres.' % fila)
    return texto


def n(texto):
    """Literal NVARCHAR solo con ASCII: cada caracter fuera de ASCII va como
    NCHAR(n) (los scripts SQL de este repo no llevan tildes, sqlcmd los lee
    con la pagina de codigos del equipo)."""
    if texto is None:
        return 'NULL'
    partes, actual = [], ''
    for c in str(texto):
        if ord(c) < 128:
            actual += c
        else:
            if actual:
                partes.append("N'" + actual.replace("'", "''") + "'")
                actual = ''
            partes.append('NCHAR(%d)' % ord(c))
    if actual or not partes:
        partes.append("N'" + actual.replace("'", "''") + "'")
    return ' + '.join(partes)


def dec(v):
    return 'NULL' if v is None else '%.2f' % v


def leer_septiembre(wb, avisos, omitidas):
    ws = wb[HOJA_SEPTIEMBRE]
    cab = [plano(ws.cell(4, c).value) for c in range(1, 8)]
    esperada = ['empresa', 'nombres apellidos', 'id sap', 'valor', 'observacion', 'fechas', 'lugar']
    if cab != esperada:
        raise Detener('la fila 4 de %s no tiene las columnas esperadas (A-G).' % HOJA_SEPTIEMBRE)
    filas = []
    for r in range(FILA_DESDE, FILA_HASTA + 1):
        empresa = str(ws.cell(r, 1).value or '').strip().upper()
        if empresa not in ('DOS', 'AGILITY'):
            raise Detener('fila %d: empresa no valida.' % r)
        valor = dinero(ws.cell(r, 4).value, r, 'VALOR')
        if valor == 0:
            omitidas.append('fila %d (%s): valor $0, no se migra (decision 7).' % (r, empresa))
            continue
        if valor < 0:
            raise Detener('fila %d: valor negativo.' % r)
        observacion = ws.cell(r, 5).value
        tipo = tipo_de(observacion)
        if tipo is None:
            raise Detener('fila %d: la observacion "%s" no corresponde a ningun tipo del catalogo.'
                          % (r, plano(observacion)))
        _, familia, sugerido, por_semana = TIPOS[tipo]
        semanas = semanas_de(observacion)
        fechas = ws.cell(r, 6).value
        if por_semana:
            if semanas is None or semanas < 1:
                raise Detener('fila %d: el tipo %s es por semana y la observacion no dice cuantas.'
                              % (r, plano(tipo)))
            lunes = lunes_de(fechas, semanas, r, avisos)
            valor_sugerido = sugerido * len(lunes)
        else:
            if semanas is not None or plano(fechas):
                raise Detener('fila %d: el tipo %s no es por semana y la fila trae semanas.'
                              % (r, plano(tipo)))
            lunes = []
            valor_sugerido = sugerido
        lugar = str(ws.cell(r, 7).value or '').strip()
        if len(lugar) > 100:
            raise Detener('fila %d: Lugar supera 100 caracteres.' % r)
        motivo = None
        if valor_sugerido is not None and valor != valor_sugerido:
            motivo = MOTIVO_AJUSTE
            avisos.append('fila %d: el valor difiere del sugerido del tipo; lleva motivo de ajuste.' % r)
        conceptos = CONCEPTO_CUMPLIMIENTO if tipo == 'Cumplimiento' else CONCEPTO
        filas.append({
            'fila': r, 'empresa': empresa,
            'ident': identificacion(empresa, ws.cell(r, 3).value, r),
            'tipo_excel': tipo, 'tipo': destino_de(tipo, valor, r), 'concepto': conceptos[empresa],
            'semanas': len(lunes), 'lunes': lunes,
            'sugerido': valor_sugerido, 'valor': valor, 'motivo': motivo, 'lugar': lugar or None,
        })
    return filas


def leer_historico(wb):
    ws = wb[HOJA_HISTORICO]
    meses = {}
    for r in range(1, ws.max_row + 1):
        anio, mes, total = ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value
        if not isinstance(anio, (int, float)) or plano(mes) not in MESES:
            continue
        clave = (int(anio), MESES[plano(mes)])
        if clave in meses:
            raise Detener('hoja %s, fila %d: el mes %04d-%02d esta repetido.' % ((HOJA_HISTORICO, r) + clave))
        meses[clave] = dinero(total, r, 'D de %s' % HOJA_HISTORICO)
    rango = [k for k in sorted(meses) if HISTORICO_DESDE <= k <= HISTORICO_HASTA]
    if len(rango) != HISTORICOS_ESPERADOS:
        raise Detener('hoja %s: hay %d meses entre %04d-%02d y %04d-%02d; se esperaban %d.'
                      % ((HOJA_HISTORICO, len(rango)) + HISTORICO_DESDE + HISTORICO_HASTA
                         + (HISTORICOS_ESPERADOS,)))
    return [(k[0], k[1], meses[k]) for k in rango], meses.get((ANIO, MES))


def armar_sql(filas, historico):
    L = []
    a = L.append
    a('/* ' + '=' * 74)
    a('   MIGRACION DEL MODULO DE BONOS - GENERADO, NO EDITAR A MANO')
    a('')
    a('   Lo produce docs/sql/generar-migracion-bonos.py desde')
    a('   RevisionBonos/9. Septiembre.xlsx.')
    a('   CONTIENE DATOS PERSONALES: codigos SAP, cedulas y montos. NO SE COMMITEA.')
    a('')
    a('   Generado el %s.' % datetime.datetime.now().strftime('%Y-%m-%d %H:%M'))
    a('   Usa los tipos que Talento Humano ya cargo en produccion (solo crea')
    a('   los que no existen y marca por semana los que hace falta).')
    a('   Re-ejecutable: el catalogo solo cambia lo que falta, septiembre se')
    a('   salta si ya tiene bonos de Origen = Migracion, y los historicos se')
    a('   insertan solo si falta su (Anio, Mes). Termina en VEREDICTO: OK.')
    a('   Si alguna fila no encuentra a su persona activa, NO migra nada y lista')
    a('   las filas del Excel por numero.')
    a('   ' + '=' * 74 + ' */')
    a('SET NOCOUNT ON;')
    a('SET QUOTED_IDENTIFIER ON;')
    a('SET ANSI_NULLS ON;')
    a('SET XACT_ABORT ON;')
    a('GO')
    a('')
    a('/* Los unicos tipos que crea la migracion; el resto ya existe en produccion. */')
    a('CREATE TABLE #Nuevo (Nombre NVARCHAR(150) NOT NULL, Familia VARCHAR(20) NOT NULL,')
    a('                     ValorSugerido DECIMAL(9,2) NULL, PorSemana BIT NOT NULL,')
    a('                     ConceptoDos VARCHAR(10) NOT NULL, ConceptoAgility VARCHAR(10) NOT NULL);')
    for nombre, familia, sugerido, por_semana in NUEVOS:
        a("INSERT INTO #Nuevo VALUES (%s, '%s', %s, %d, '%s', '%s');"
          % (n(nombre), familia, dec(sugerido), 1 if por_semana else 0,
             CONCEPTO['DOS'], CONCEPTO['AGILITY']))
    a('')
    a('/* Tipos de produccion no semanales que se pagan por semana. */')
    a('CREATE TABLE #Semanal (Nombre NVARCHAR(150) NOT NULL PRIMARY KEY);')
    for nombre in SEMANALES:
        a('INSERT INTO #Semanal VALUES (%s);' % n(nombre))
    a('')
    a('/* Una fila por fila del Excel (sin la de $0). Ident: DOS = Cod_Sap,')
    a('   AGILITY = cedula rellenada a 10. Tipo = el de produccion; TipoExcel =')
    a('   como lo llamaba el Excel. Concepto = el que se pago en septiembre. */')
    a('CREATE TABLE #X (Fila INT NOT NULL PRIMARY KEY, Empresa VARCHAR(10) NOT NULL,')
    a('                 Ident VARCHAR(32) NOT NULL, Tipo NVARCHAR(150) NOT NULL,')
    a('                 TipoExcel NVARCHAR(150) NOT NULL, Concepto VARCHAR(10) NOT NULL, Semanas INT NOT NULL,')
    a('                 ValorSugerido DECIMAL(9,2) NULL, Valor DECIMAL(9,2) NOT NULL,')
    a('                 MotivoAjuste NVARCHAR(500) NULL, Lugar NVARCHAR(100) NULL);')
    for f in filas:
        a("INSERT INTO #X VALUES (%d, '%s', '%s', %s, %s, '%s', %d, %s, %s, %s, %s);"
          % (f['fila'], f['empresa'], f['ident'], n(f['tipo']), n(f['tipo_excel']), f['concepto'],
             f['semanas'], dec(f['sugerido']), dec(f['valor']), n(f['motivo']), n(f['lugar'])))
    a('')
    a('CREATE TABLE #W (Fila INT NOT NULL, Lunes DATE NOT NULL, PRIMARY KEY (Fila, Lunes));')
    for f in filas:
        for lun in f['lunes']:
            a("INSERT INTO #W VALUES (%d, '%s');" % (f['fila'], lun.strftime('%Y%m%d')))
    a('')
    a('CREATE TABLE #H (Anio INT NOT NULL, Mes INT NOT NULL, Total DECIMAL(11,2) NOT NULL,')
    a('                 PRIMARY KEY (Anio, Mes));')
    for anio, mes, total in historico:
        a('INSERT INTO #H VALUES (%d, %d, %s);' % (anio, mes, dec(total)))
    a('')
    a('/* Personas activas: el UNICO criterio es Sp_RTA_BonoPersonasActivas. */')
    a('/* Columnas holgadas: un truncamiento aqui trae el valor en el mensaje. */')
    a('CREATE TABLE #P (CodUsuario VARCHAR(100), Nombre NVARCHAR(500), Correo NVARCHAR(500),')
    a('                 CodJefe VARCHAR(100), Empresa VARCHAR(100), CodSap VARCHAR(100),')
    a('                 Cedula VARCHAR(100), IdPerfil INT);')
    a('INSERT INTO #P EXEC dbo.Sp_RTA_BonoPersonasActivas;')
    a('GO')
    a('')
    a('BEGIN TRY')
    a('BEGIN TRANSACTION;')
    a('')
    a('DECLARE @Ya INT = (SELECT COUNT(*) FROM dbo.Bono_Detalle d')
    a('                     JOIN dbo.Bono_Periodo p ON p.IdPeriodo = d.IdPeriodo')
    a("                    WHERE p.Anio = %d AND p.Mes = %d AND d.Origen = 'Migracion');" % (ANIO, MES))
    a('')
    a('IF @Ya = 0')
    a('BEGIN')
    a('    /* Cada fila con exactamente una persona activa, de la misma empresa. */')
    a('    SELECT x.Fila, x.Empresa,')
    a('           Encontradas = (SELECT COUNT(DISTINCT p.CodUsuario) FROM #P p')
    a("                           WHERE (x.Empresa = 'DOS' AND p.CodSap = x.Ident)")
    a("                              OR (x.Empresa = 'AGILITY' AND p.Cedula <> ''")
    a("                                  AND RIGHT('0000000000' + p.Cedula, 10) = x.Ident)),")
    a('           MismaEmpresa = (SELECT COUNT(DISTINCT p.CodUsuario) FROM #P p')
    a('                            WHERE UPPER(p.Empresa) = x.Empresa')
    a("                              AND ((x.Empresa = 'DOS' AND p.CodSap = x.Ident)")
    a("                                OR (x.Empresa = 'AGILITY' AND p.Cedula <> ''")
    a("                                    AND RIGHT('0000000000' + p.Cedula, 10) = x.Ident)))")
    a('      INTO #Busqueda')
    a('      FROM #X x;')
    a('')
    a('    IF EXISTS (SELECT 1 FROM #Busqueda WHERE Encontradas <> 1 OR MismaEmpresa <> 1)')
    a('    BEGIN')
    a("        PRINT 'Filas del Excel sin una persona activa unica (fila, empresa, motivo):';")
    a('        DECLARE @Lista NVARCHAR(MAX) = N\'\';')
    a("        SELECT @Lista = @Lista + N'  fila ' + CONVERT(NVARCHAR(10), Fila) + N' (' + Empresa + N'): '")
    a("             + CASE WHEN Encontradas = 0 THEN N'no aparece entre las personas activas'")
    a("                    WHEN Encontradas > 1 THEN N'aparece en mas de una persona activa'")
    a("                    ELSE N'la persona activa tiene otra empresa' END + NCHAR(10)")
    a('          FROM #Busqueda WHERE Encontradas <> 1 OR MismaEmpresa <> 1 ORDER BY Fila;')
    a('        PRINT @Lista;')
    a("        SELECT Fila, Empresa, Encontradas, MismaEmpresa FROM #Busqueda")
    a('         WHERE Encontradas <> 1 OR MismaEmpresa <> 1 ORDER BY Fila;')
    a("        RAISERROR('FALLO: hay filas del Excel sin su persona. No se migro nada (ni el catalogo).', 16, 1);")
    a('    END')
    a('')
    a("    IF EXISTS (SELECT 1 FROM dbo.Bono_Periodo WHERE Anio = %d AND Mes = %d)" % (ANIO, MES))
    a("        RAISERROR('FALLO: septiembre 2026 ya existe y no tiene bonos de la migracion. No se migro nada.', 16, 1);")
    a('END')
    a('')
    a('/* ---------------------------------------------- 1. catalogo ----------- */')
    a('/* El catalogo es de Talento Humano: no se cambian nombres, valores ni')
    a('   conceptos. Solo se crean los tipos de #Nuevo y se marcan semanales los')
    a('   de #Semanal; cada cambio queda en la bitacora como los del Catalogo. */')
    a('DECLARE @Creados TABLE (IdTipo INT NOT NULL, Nombre NVARCHAR(150) NOT NULL);')
    a('INSERT INTO dbo.Bono_Tipo (Nombre, Familia, ValorSugerido, PorSemana, Activo, Usuario_Registro)')
    a('OUTPUT inserted.IdTipo, inserted.Nombre INTO @Creados (IdTipo, Nombre)')
    a("SELECT o.Nombre, o.Familia, o.ValorSugerido, o.PorSemana, 1, '%s'" % USUARIO)
    a('  FROM #Nuevo o')
    a(' WHERE NOT EXISTS (SELECT 1 FROM dbo.Bono_Tipo t WHERE t.Nombre = o.Nombre);')
    a('')
    a('INSERT INTO dbo.Bono_TipoConcepto (IdTipo, Empresa, CodigoConcepto, Usuario_Registro)')
    a("SELECT t.IdTipo, c.Empresa, c.CodigoConcepto, '%s'" % USUARIO)
    a('  FROM #Nuevo o')
    a('  JOIN @Creados t ON t.Nombre = o.Nombre')
    a("  CROSS APPLY (VALUES ('DOS', o.ConceptoDos), ('AGILITY', o.ConceptoAgility)) AS c (Empresa, CodigoConcepto);")
    a('')
    a('INSERT INTO dbo.Bono_Bitacora (IdBono, IdPeriodo, Accion, ValorAnterior, ValorNuevo, Detalle, Usuario)')
    a("SELECT NULL, NULL, 'Catalogo', NULL, LEFT(c.Nombre, 200),")
    a("       N'Crear tipo ' + CONVERT(NVARCHAR(12), c.IdTipo) + N' desde la migracion de septiembre', '%s'" % USUARIO)
    a('  FROM @Creados c;')
    a('')
    a('DECLARE @Marcados TABLE (IdTipo INT NOT NULL, Nombre NVARCHAR(150) NOT NULL);')
    a("UPDATE t SET PorSemana = 1, Usuario_Modificacion = '%s', Fecha_Modificacion = GETDATE()" % USUARIO)
    a('OUTPUT inserted.IdTipo, inserted.Nombre INTO @Marcados (IdTipo, Nombre)')
    a('  FROM dbo.Bono_Tipo t')
    a('  JOIN #Semanal s ON s.Nombre = t.Nombre')
    a(' WHERE t.PorSemana = 0;')
    a('')
    a('INSERT INTO dbo.Bono_Bitacora (IdBono, IdPeriodo, Accion, ValorAnterior, ValorNuevo, Detalle, Usuario)')
    a("SELECT NULL, NULL, 'Catalogo', N'no semanal', N'por semana',")
    a("       N'Tipo ' + CONVERT(NVARCHAR(12), m.IdTipo) + N' (' + LEFT(m.Nombre, 150)")
    a("       + N') pasa a por semana: sus bonos de septiembre traen semanas', '%s'" % USUARIO)
    a('  FROM @Marcados m;')
    a('')
    a('/* ---------------------------------------------- 2. septiembre --------- */')
    a('IF @Ya = 0')
    a('BEGIN')
    a('    /* Cada tipo destino existe en produccion, una sola vez. */')
    a('    IF EXISTS (SELECT 1 FROM (SELECT DISTINCT Tipo FROM #X) x')
    a('                WHERE (SELECT COUNT(*) FROM dbo.Bono_Tipo t WHERE t.Nombre = x.Tipo) <> 1)')
    a('    BEGIN')
    a("        PRINT 'Tipos de destino que no existen (o estan repetidos) en el catalogo:';")
    a('        DECLARE @Faltan NVARCHAR(MAX) = N\'\';')
    a("        SELECT @Faltan = @Faltan + N'  ' + x.Tipo + NCHAR(10)")
    a('          FROM (SELECT DISTINCT Tipo FROM #X) x')
    a('         WHERE (SELECT COUNT(*) FROM dbo.Bono_Tipo t WHERE t.Nombre = x.Tipo) <> 1;')
    a('        PRINT @Faltan;')
    a("        RAISERROR('FALLO: falta un tipo de destino en el catalogo. No se migro nada.', 16, 1);")
    a('    END')
    a('')
    a('    /* Una fila con semanas solo va a un tipo por semana, y al reves. */')
    a('    IF EXISTS (SELECT 1 FROM #X x JOIN dbo.Bono_Tipo t ON t.Nombre = x.Tipo')
    a('                WHERE t.PorSemana <> CASE WHEN x.Semanas > 0 THEN 1 ELSE 0 END)')
    a('    BEGIN')
    a("        PRINT 'Filas cuyo tipo de destino no coincide en por semana (fila, tipo):';")
    a('        DECLARE @NoCoinciden NVARCHAR(MAX) = N\'\';')
    a("        SELECT @NoCoinciden = @NoCoinciden + N'  fila ' + CONVERT(NVARCHAR(10), x.Fila) + N': ' + x.Tipo + NCHAR(10)")
    a('          FROM #X x JOIN dbo.Bono_Tipo t ON t.Nombre = x.Tipo')
    a('         WHERE t.PorSemana <> CASE WHEN x.Semanas > 0 THEN 1 ELSE 0 END ORDER BY x.Fila;')
    a('        PRINT @NoCoinciden;')
    a("        RAISERROR('FALLO: hay filas con un tipo de destino que no coincide en por semana. No se migro nada.', 16, 1);")
    a('    END')
    a('')
    a('    INSERT INTO dbo.Bono_Periodo (Anio, Mes, Estado, FechaApertura, UsuarioApertura,')
    a('                                  FechaCierre, UsuarioCierre, Usuario_Registro)')
    a("    VALUES (%d, %d, 'Cerrado', GETDATE(), '%s', GETDATE(), '%s', '%s');" % (ANIO, MES, USUARIO, USUARIO, USUARIO))
    a('    DECLARE @IdPeriodo INT = SCOPE_IDENTITY();')
    a('')
    a('    CREATE TABLE #Mapa (Fila INT NOT NULL PRIMARY KEY, IdBono INT NOT NULL);')
    a('')
    a('    /* MERGE con ON 1 = 0 solo para que OUTPUT pueda devolver la fila del')
    a('       Excel junto al IdBono nuevo (INSERT ... OUTPUT no ve el origen). */')
    a('    MERGE dbo.Bono_Detalle AS d')
    a('    USING (SELECT x.Fila, p.CodUsuario, Empresa = UPPER(LTRIM(RTRIM(p.Empresa))), p.Nombre,')
    a("                  Identificacion = CASE WHEN x.Empresa = 'DOS' THEN p.CodSap")
    a("                                        ELSE RIGHT('0000000000' + p.Cedula, 10) END,")
    a('                  x.Lugar, t.IdTipo, x.Semanas, x.ValorSugerido, x.Valor, x.MotivoAjuste,')
    a("                  CodJefe = NULLIF(p.CodJefe, ''), NombreTipo = t.Nombre, CodigoConcepto = x.Concepto")
    a('             FROM #X x')
    a("             JOIN #P p ON (x.Empresa = 'DOS' AND p.CodSap = x.Ident)")
    a("                       OR (x.Empresa = 'AGILITY' AND p.Cedula <> ''")
    a("                           AND RIGHT('0000000000' + p.Cedula, 10) = x.Ident)")
    a('             JOIN dbo.Bono_Tipo t ON t.Nombre = x.Tipo) AS o')
    a('       ON 1 = 0')
    a('     WHEN NOT MATCHED THEN')
    a('        INSERT (IdPeriodo, CodUsuario, Empresa, Nombre, Identificacion, Lugar, IdTipo, Semanas,')
    a('                ValorSugerido, Valor, MotivoAjuste, Estado, Origen, CodJefe, UsuarioRevisa,')
    a('                FechaRevisa, NombreTipo, CodigoConcepto, Usuario_Registro)')
    a('        VALUES (@IdPeriodo, o.CodUsuario, o.Empresa, o.Nombre, o.Identificacion, o.Lugar, o.IdTipo,')
    a("                o.Semanas, o.ValorSugerido, o.Valor, o.MotivoAjuste, 'Aprobado', 'Migracion',")
    a("                o.CodJefe, '%s', GETDATE(), o.NombreTipo, o.CodigoConcepto, '%s')" % (USUARIO, USUARIO))
    a('    OUTPUT o.Fila, inserted.IdBono INTO #Mapa (Fila, IdBono);')
    a('')
    a('    IF (SELECT COUNT(*) FROM #Mapa) <> (SELECT COUNT(*) FROM #X)')
    a("        RAISERROR('FALLO: no se insertaron todas las filas del Excel. No se migro nada.', 16, 1);")
    a('')
    a('    INSERT INTO dbo.Bono_DetalleSemana (IdBono, Lunes, Usuario_Registro)')
    a("    SELECT m.IdBono, w.Lunes, '%s' FROM #W w JOIN #Mapa m ON m.Fila = w.Fila;" % USUARIO)
    a('')
    a('    /* Ya hay bonos cargados desde la pantalla: una semana de septiembre que')
    a('       tambien pague uno de ellos (misma persona y tipo) revierte todo. */')
    a('    IF EXISTS (SELECT 1 FROM dbo.Bono_DetalleSemana s')
    a('                 JOIN dbo.Bono_Detalle d ON d.IdBono = s.IdBono')
    a("                WHERE d.Estado <> 'Rechazado'")
    a('                GROUP BY d.CodUsuario, d.IdTipo, s.Lunes HAVING COUNT(*) > 1)')
    a("        RAISERROR('FALLO: una semana de septiembre ya la paga otro bono de la misma persona y tipo. No se migro nada.', 16, 1);")
    a('')
    a('    INSERT INTO dbo.Bono_Bitacora (IdBono, IdPeriodo, Accion, ValorNuevo, Detalle, Usuario)')
    a("    SELECT m.IdBono, @IdPeriodo, 'Crear', 'Aprobado',")
    a("           N'Migracion desde el Excel de septiembre, fila ' + CONVERT(NVARCHAR(10), m.Fila)")
    a("           + N' (en el Excel: ' + x.TipoExcel + N')', '%s'" % USUARIO)
    a('      FROM #Mapa m JOIN #X x ON x.Fila = m.Fila;')
    a('    INSERT INTO dbo.Bono_Bitacora (IdPeriodo, Accion, ValorNuevo, Detalle, Usuario)')
    a("    VALUES (@IdPeriodo, 'Cerrar', 'Cerrado', N'Migracion desde el Excel de septiembre', '%s');" % USUARIO)
    a('    DECLARE @Migrados INT = (SELECT COUNT(*) FROM #Mapa);')
    a("    PRINT 'Septiembre 2026 migrado: ' + CONVERT(VARCHAR(10), @Migrados) + ' bonos.';")
    a('END')
    a('ELSE')
    a("    PRINT 'Septiembre 2026 ya estaba migrado: no se toca.';")
    a('')
    a('/* ---------------------------------------------- 3. historicos --------- */')
    a('INSERT INTO dbo.Bono_Periodo (Anio, Mes, Estado, TotalHistorico, Usuario_Registro)')
    a("SELECT h.Anio, h.Mes, 'Historico', h.Total, '%s'" % USUARIO)
    a('  FROM #H h')
    a(' WHERE NOT EXISTS (SELECT 1 FROM dbo.Bono_Periodo p WHERE p.Anio = h.Anio AND p.Mes = h.Mes);')
    a("PRINT 'Historicos agregados: ' + CONVERT(VARCHAR(10), @@ROWCOUNT) + '.';")
    a('')
    a('COMMIT TRANSACTION;')
    a('END TRY')
    a('BEGIN CATCH')
    a('    IF @@TRANCOUNT > 0 ROLLBACK TRANSACTION;')
    a("    PRINT 'Error, la migracion quedo revertida. ERROR_NUMBER=' + CONVERT(VARCHAR(12), ERROR_NUMBER())")
    a("        + ' ERROR_LINE=' + CONVERT(VARCHAR(12), ERROR_LINE());")
    a('    /* 2628 y 8152 (truncamiento) traen el valor dentro del mensaje. */')
    a('    IF ERROR_NUMBER() NOT IN (2628, 8152)')
    a("        PRINT 'Detalle: ' + ERROR_MESSAGE();")
    a('END CATCH')
    a('GO')
    a('')
    a('/* ------------------------------------------------- veredicto -------------- */')
    a('/* Solo lectura. Sin GO entre el DECLARE de @Fallos y su ultimo uso. */')
    a('SET NOCOUNT ON;')
    a('DECLARE @Fallos INT = 0;')
    a('')
    a('DECLARE @Sep DECIMAL(11,2) = (SELECT ISNULL(SUM(d.Valor), 0) FROM dbo.Bono_Detalle d')
    a('                                JOIN dbo.Bono_Periodo p ON p.IdPeriodo = d.IdPeriodo')
    a("                               WHERE p.Anio = %d AND p.Mes = %d AND d.Estado = 'Aprobado');" % (ANIO, MES))
    a('IF @Sep <> %s' % dec(TOTAL_SEPTIEMBRE))
    a('BEGIN')
    a("    PRINT 'FALLO: septiembre 2026 suma ' + CONVERT(VARCHAR(20), @Sep) + ' y debe sumar %s.';" % dec(TOTAL_SEPTIEMBRE))
    a('    SET @Fallos += 1;')
    a('END')
    a('')
    a('DECLARE @BonosSep INT = (SELECT COUNT(*) FROM dbo.Bono_Detalle d')
    a('                           JOIN dbo.Bono_Periodo p ON p.IdPeriodo = d.IdPeriodo')
    a("                          WHERE p.Anio = %d AND p.Mes = %d AND d.Origen = 'Migracion');" % (ANIO, MES))
    a('IF @BonosSep <> %d' % len(filas))
    a('BEGIN')
    a("    PRINT 'FALLO: septiembre 2026 tiene ' + CONVERT(VARCHAR(10), @BonosSep) + ' bonos migrados y debe tener %d.';" % len(filas))
    a('    SET @Fallos += 1;')
    a('END')
    a('')
    a("IF NOT EXISTS (SELECT 1 FROM dbo.Bono_Periodo WHERE Anio = %d AND Mes = %d AND Estado = 'Cerrado')" % (ANIO, MES))
    a('BEGIN')
    a("    PRINT 'FALLO: septiembre 2026 no existe como periodo Cerrado.';")
    a('    SET @Fallos += 1;')
    a('END')
    a('')
    a("DECLARE @Hist INT = (SELECT COUNT(*) FROM dbo.Bono_Periodo WHERE Estado = 'Historico'")
    a('                        AND Anio * 100 + Mes BETWEEN %d AND %d' % (HISTORICO_DESDE[0] * 100 + HISTORICO_DESDE[1],
                                                                       HISTORICO_HASTA[0] * 100 + HISTORICO_HASTA[1]))
    a('                        AND TotalHistorico IS NOT NULL);')
    a('IF @Hist <> %d' % HISTORICOS_ESPERADOS)
    a('BEGIN')
    a("    PRINT 'FALLO: hay ' + CONVERT(VARCHAR(10), @Hist) + ' periodos Historico con total; se esperaban %d.';" % HISTORICOS_ESPERADOS)
    a('    SET @Fallos += 1;')
    a('END')
    a('')
    a('/* Ningun doble pago: misma persona, mismo tipo, mismo lunes, dos bonos no rechazados. */')
    a('DECLARE @Doble INT = (SELECT COUNT(*) FROM (')
    a('    SELECT d.CodUsuario, d.IdTipo, s.Lunes FROM dbo.Bono_DetalleSemana s')
    a('      JOIN dbo.Bono_Detalle d ON d.IdBono = s.IdBono')
    a("     WHERE d.Estado <> 'Rechazado'")
    a('     GROUP BY d.CodUsuario, d.IdTipo, s.Lunes HAVING COUNT(*) > 1) AS x);')
    a('IF @Doble <> 0')
    a('BEGIN')
    a("    PRINT 'FALLO: ' + CONVERT(VARCHAR(10), @Doble) + ' semanas pagadas dos veces.';")
    a('    SET @Fallos += 1;')
    a('END')
    a('')
    a('/* Cada bono migrado tiene tantos lunes como Semanas, todos desde el 2026-08-03. */')
    a('DECLARE @Semanas INT = (SELECT COUNT(*) FROM dbo.Bono_Detalle d')
    a("                          WHERE d.Origen = 'Migracion'")
    a('                            AND (d.Semanas <> (SELECT COUNT(*) FROM dbo.Bono_DetalleSemana s WHERE s.IdBono = d.IdBono)')
    a('                                 OR EXISTS (SELECT 1 FROM dbo.Bono_DetalleSemana s')
    a("                                             WHERE s.IdBono = d.IdBono AND s.Lunes < '%s')));"
      % PRIMER_LUNES.strftime('%Y%m%d'))
    a('IF @Semanas <> 0')
    a('BEGIN')
    a("    PRINT 'FALLO: ' + CONVERT(VARCHAR(10), @Semanas) + ' bonos migrados con semanas que no cuadran.';")
    a('    SET @Fallos += 1;')
    a('END')
    a('')
    a('/* Los bonos migrados traen sus copias completas. */')
    a('DECLARE @Copias INT = (SELECT COUNT(*) FROM dbo.Bono_Detalle')
    a("                         WHERE Origen = 'Migracion'")
    a("                           AND (ISNULL(Empresa, '') NOT IN ('DOS', 'AGILITY') OR ISNULL(Nombre, N'') = N''")
    a("                                OR ISNULL(Identificacion, '') = '' OR ISNULL(NombreTipo, N'') = N''")
    a("                                OR ISNULL(CodigoConcepto, '') = ''));")
    a('IF @Copias <> 0')
    a('BEGIN')
    a("    PRINT 'FALLO: ' + CONVERT(VARCHAR(10), @Copias) + ' bonos migrados sin empresa, nombre, identificacion, tipo o concepto.';")
    a('    SET @Fallos += 1;')
    a('END')
    a('')
    a('/* Cada bono migrado va al tipo de produccion que le toca. */')
    a('DECLARE @MalTipo INT = (SELECT COUNT(*) FROM dbo.Bono_Detalle d')
    a('                          JOIN dbo.Bono_Periodo p ON p.IdPeriodo = d.IdPeriodo')
    a('                          JOIN dbo.Bono_Tipo t ON t.IdTipo = d.IdTipo')
    a("                         WHERE p.Anio = %d AND p.Mes = %d AND d.Origen = 'Migracion'" % (ANIO, MES))
    a('                           AND t.Nombre NOT IN (SELECT Tipo FROM #X));')
    a('IF @MalTipo <> 0')
    a('BEGIN')
    a("    PRINT 'FALLO: ' + CONVERT(VARCHAR(10), @MalTipo) + ' bonos migrados con un tipo fuera de la correspondencia.';")
    a('    SET @Fallos += 1;')
    a('END')
    a('')
    a('DECLARE @NoSemanal INT = (SELECT COUNT(*) FROM #Semanal s')
    a('                           WHERE NOT EXISTS (SELECT 1 FROM dbo.Bono_Tipo t WHERE t.Nombre = s.Nombre AND t.PorSemana = 1));')
    a('IF @NoSemanal <> 0')
    a('BEGIN')
    a("    PRINT 'FALLO: ' + CONVERT(VARCHAR(10), @NoSemanal) + ' tipos que debian quedar por semana no lo estan.';")
    a('    SET @Fallos += 1;')
    a('END')
    a('')
    a('/* Aviso, no falla: los conceptos del catalogo los completa Talento Humano. */')
    a('DECLARE @SinConcepto INT = (SELECT COUNT(*) FROM dbo.Bono_Tipo t')
    a('                              WHERE t.Activo = 1')
    a("                                AND (SELECT COUNT(*) FROM dbo.Bono_TipoConcepto k WHERE k.IdTipo = t.IdTipo")
    a("                                        AND k.Empresa IN ('DOS', 'AGILITY')) <> 2);")
    a('IF @SinConcepto <> 0')
    a("    PRINT 'AVISO: ' + CONVERT(VARCHAR(10), @SinConcepto) + ' tipos activos sin concepto DOS o AGILITY (los completa Talento Humano en el Catalogo).';")
    a('')
    a('SELECT TotalSeptiembre = @Sep, BonosSeptiembre = @BonosSep, Historicos = @Hist,')
    a('       DoblePago = @Doble, Fallos = @Fallos;')
    a('IF @Fallos = 0')
    a("    PRINT 'VEREDICTO: OK';")
    a('ELSE')
    a("    PRINT 'VEREDICTO: FALLA (' + CONVERT(VARCHAR(10), @Fallos) + ' comprobaciones). Revisar arriba.';")
    a("SELECT Veredicto = CASE WHEN @Fallos = 0 THEN 'VEREDICTO: OK' ELSE 'VEREDICTO: FALLA' END;")
    a('GO')
    return L


def main():
    if not os.path.exists(EXCEL):
        print('No se encontro el Excel en %s' % EXCEL)
        return 1

    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    avisos, omitidas = [], []
    try:
        filas = leer_septiembre(wb, avisos, omitidas)
        historico, total_vista = leer_historico(wb)

        total = sum((f['valor'] for f in filas), Decimal('0.00'))
        if total != TOTAL_SEPTIEMBRE:
            raise Detener('las filas de septiembre suman %s y deben sumar %s.' % (total, TOTAL_SEPTIEMBRE))
        if total_vista is not None and total_vista != TOTAL_SEPTIEMBRE:
            raise Detener('%s dice que septiembre suma %s, no %s.' % (HOJA_HISTORICO, total_vista, TOTAL_SEPTIEMBRE))

        # Doble pago dentro del propio Excel: misma persona, tipo y lunes en
        # dos filas. El script lo vuelve a comprobar en la base.
        semanas = Counter((f['empresa'], f['ident'], f['tipo'], lun) for f in filas for lun in f['lunes'])
        dobles = sorted(set(f['fila'] for f in filas for lun in f['lunes']
                            if semanas[(f['empresa'], f['ident'], f['tipo'], lun)] > 1))
        if dobles:
            raise Detener('doble pago dentro del Excel, filas: %s.' % ', '.join(str(x) for x in dobles))
    except Detener as e:
        print('No se genera nada: %s' % e)
        return 1

    L = armar_sql(filas, historico)
    texto = '\n'.join(L) + '\n'
    try:
        texto.encode('ascii')
    except UnicodeEncodeError:
        print('No se genera nada: el script saldria con caracteres fuera de ASCII.')
        return 1

    if not os.path.isdir(SALIDA_DIR):
        os.makedirs(SALIDA_DIR)
    # utf-8-sig como el generador de horas extras; el contenido es ASCII
    # (las tildes van como NCHAR), asi que el BOM es solo por coherencia.
    with open(SALIDA, 'w', encoding='utf-8-sig', newline='\n') as f:
        f.write(texto)

    print('Generado: %s' % SALIDA)
    print('  bonos de septiembre: %d (suman %s)' % (len(filas), total))
    print('  semanas:             %d' % sum(f['semanas'] for f in filas))
    print('  historicos:          %d' % len(historico))
    print('  tipos de destino:    %d (crea %d, marca por semana %d)'
          % (len(set(f['tipo'] for f in filas)), len(NUEVOS), len(SEMANALES)))
    print('Correspondencia (Excel -> produccion: filas):')
    for (excel, prod), veces in sorted(Counter((f['tipo_excel'], f['tipo']) for f in filas).items()):
        print('  %s -> %s: %d' % (plano(excel), plano(prod), veces))
    print('Omitidas (%d):' % len(omitidas))
    for x in omitidas:
        print('  - %s' % x)
    print('Ajustes (%d):' % len(avisos))
    for x in avisos:
        print('  - %s' % x)
    print('NO COMMITEAR ese archivo: contiene datos personales.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
