# -*- coding: utf-8 -*-
"""
Genera la carga del motor de vacaciones y su prueba de oro desde el Excel de GTH.

POR QUE EXISTE ESTE ARCHIVO Y NO UN .sql CON LOS DATOS DENTRO:
el Excel trae nombres, cedulas y sueldos, y este repositorio es PUBLICO. Se
versiona la logica de la carga -este archivo- y no su resultado.

Uso:
    python -I docs/sql/generar-carga-vacaciones.py "SISTEMA_VACACIONES_DOS (15).xlsx"

Escribe en docs/sql/carga-generada/ (en .gitignore):
  vacaciones-carga.sql       Vac_Colaborador, Vac_Movimiento (Origen EXCEL) y
                             Vac_FeriadoHistorico (2005-2025). Reemplaza la
                             carga anterior; se niega si ya hay movimientos
                             Origen SISTEMA.
  vacaciones-prueba-oro.sql  Lo que el Excel calculo (hojas SALDOS,
                             CALCULO_PERIODOS y REGISTRO) contra lo que
                             devuelve Sp_RTA_VacCalcular a la misma fecha.
                             Imprime PASS / FAIL y VEREDICTO.
  vacaciones-sueldos.sql     Vac_Sueldo (Origen EXCEL), por cedula, desde
                             COLABORADORES!J. Reemplaza los de Origen EXCEL;
                             no toca los cargados desde el sistema. La nomina
                             usa HE_Salario cuando la persona lo tiene.

La prueba usa los valores que el Excel guardo la ultima vez que se calculo, asi
que el archivo debe haberse abierto y guardado en Excel (no vale uno generado
por otro programa). La fecha de calculo es la que quedo en PARAMETROS!B5.

Lo que imprime este generador lleva numero de fila, nunca nombre ni cedula.
"""
import datetime
import os
import sys

import openpyxl

FILA_INICIO = 5
ESTADOS = {
    'APROBADO': 'APROBADO', 'POR APROBAR': 'POR APROBAR', 'DADA BAJA': 'DADA BAJA',
    'HISTÓRICO': 'HISTORICO', 'HISTORICO': 'HISTORICO', 'PROCESADO': 'PROCESADO',
    'REVISAR': 'REVISAR',
}
TIPOS = ('VACACIONES', 'PERMISO POR HORAS', 'COMPRA DE VACACIONES')
# Situacion del Excel (REGISTRO!M) -> la que devuelve el motor
SITUACIONES = {
    'SIN FECHA - REVISAR': 'SIN FECHA',
    'ANULADO (no descuenta)': 'ANULADO',
    'RELACIÓN ANTERIOR (no cuenta)': 'RELACION ANTERIOR',
    'ANTES DEL CORTE (en saldo real)': 'ANTES DEL CORTE',
    'POSTERIOR A LA FECHA DE CÁLCULO': 'POSTERIOR',
}
ESTADOS_SALDO = {'AL DÍA': 'AL DIA', 'ACUMULANDO': 'ACUMULANDO', 'POR VENCER': 'POR VENCER', 'ANTICIPO': 'ANTICIPO'}


def txt(v):
    if v is None:
        return 'NULL'
    return "N'" + str(v).strip().replace("'", "''") + "'"


def fecha(v):
    if v is None or v == '':
        return 'NULL'
    if isinstance(v, datetime.datetime):
        v = v.date()
    return "'" + v.strftime('%Y%m%d') + "'"


def num(v):
    if v is None or v == '':
        return 'NULL'
    return repr(float(v))


def main():
    if len(sys.argv) != 2:
        sys.exit('Uso: python -I docs/sql/generar-carga-vacaciones.py <excel>')
    libro = openpyxl.load_workbook(sys.argv[1], data_only=True)
    salida = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'carga-generada')
    os.makedirs(salida, exist_ok=True)

    f_rep = libro['PARAMETROS']['B5'].value
    if not isinstance(f_rep, datetime.datetime):
        sys.exit('PARAMETROS!B5 no tiene fecha calculada: abrir y guardar el Excel en Excel.')

    # ------------------------------------------------------------ colaboradores
    colabs = []          # (n, cedula, nombre, ...)
    por_nombre = {}
    for i, r in enumerate(libro['COLABORADORES'].iter_rows(min_row=FILA_INICIO, values_only=True)):
        if not r[2]:
            continue
        n = i + 1                         # el N de la hoja: posicion dentro del rango C5:C154
        if r[1] is None:
            sys.exit('COLABORADORES fila %d: sin cedula.' % (i + FILA_INICIO))
        if r[6] is None:
            sys.exit('COLABORADORES fila %d: sin fecha de ingreso.' % (i + FILA_INICIO))
        corte = r[10] if r[10] else None
        saldo_corte = r[11] if corte else None
        colabs.append((n, str(r[1]).strip(), r[2], r[3], r[4], r[5], r[6], (r[7] or 'ACTIVO'), r[8],
                       corte, saldo_corte, r[12]))
        por_nombre[r[2]] = n

    # ------------------------------------------------------------- movimientos
    movs = []
    for i, r in enumerate(libro['REGISTRO'].iter_rows(min_row=FILA_INICIO, values_only=True)):
        if not r[0]:
            continue
        fila = i + FILA_INICIO
        if r[0] not in por_nombre:
            sys.exit('REGISTRO fila %d: la persona no esta en COLABORADORES.' % fila)
        if r[1] not in TIPOS:
            sys.exit('REGISTRO fila %d: tipo desconocido.' % fila)
        if r[2] not in ESTADOS:
            sys.exit('REGISTRO fila %d: estado desconocido.' % fila)
        beneficio = 1 if r[8] and 'BENEFICIO DOS' in str(r[8]).upper() else 0
        movs.append((fila, por_nombre[r[0]], r[1], ESTADOS[r[2]], r[3], r[4], r[5], r[6], r[7], beneficio, r[8],
                     r[11], r[12], r[25]))

    # --------------------------------------------------------------- feriados
    feriados = []
    for r in libro['PARAMETROS'].iter_rows(min_row=FILA_INICIO, values_only=True):
        if isinstance(r[7], datetime.datetime) and r[7].year < 2026:
            feriados.append((r[7], r[9] or 'Feriado'))

    # ------------------------------------------------------------- carga.sql
    c = ["/* Generado por docs/sql/generar-carga-vacaciones.py el %s. NO VERSIONAR: datos personales. */"
         % datetime.date.today().isoformat(),
         'SET NOCOUNT ON;', 'SET XACT_ABORT ON;', 'SET QUOTED_IDENTIFIER ON;',
         "IF EXISTS (SELECT 1 FROM dbo.Vac_Movimiento WHERE Origen = 'SISTEMA')",
         "BEGIN RAISERROR('Ya hay movimientos Origen SISTEMA: la carga no se puede reemplazar.', 16, 1); RETURN; END",
         # Los folios guardan IdColaborador, y la carga vuelve a crear las personas con Id nuevos:
         # un folio quedaria apuntando a otra persona.
         "IF OBJECT_ID('dbo.Vac_EstadoEmitido', 'U') IS NOT NULL AND EXISTS (SELECT 1 FROM dbo.Vac_EstadoEmitido)",
         "BEGIN RAISERROR('Ya hay estados individuales emitidos (folios): la carga no se puede reemplazar.', 16, 1); RETURN; END",
         # GTH ya corrigio cortes desde la pantalla de ajustes: la carga los pisaria con los del Excel.
         "IF OBJECT_ID('dbo.Vac_Bitacora', 'U') IS NOT NULL AND EXISTS (SELECT 1 FROM dbo.Vac_Bitacora)",
         "BEGIN RAISERROR('Ya hay ajustes de GTH en Vac_Bitacora: la carga no se puede reemplazar.', 16, 1); RETURN; END",
         'BEGIN TRANSACTION;',
         "DELETE FROM dbo.Vac_Movimiento WHERE Origen = 'EXCEL';",
         'DELETE FROM dbo.Vac_Colaborador;',
         "DELETE FROM dbo.Vac_FeriadoHistorico;"]
    for (n, ced, nom, emp, area, cargo, ing, est, sal, corte, sc, obs) in colabs:
        c.append('INSERT INTO dbo.Vac_Colaborador (Cedula, Nombre, Empresa, Area, Cargo, FechaIngreso, Estado, '
                 'FechaSalida, FechaCorte, SaldoCorte, Observacion, FilaExcel, Usuario_Registro) VALUES '
                 "(%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %d, 'CARGA_EXCEL');"
                 % (txt(ced), txt(nom), txt(emp), txt(area), txt(cargo), fecha(ing), txt(est), fecha(sal),
                    fecha(corte), num(sc), txt(obs), n))
    for (fila, n, tipo, est, d, h, dias, horas, mins, ben, obs, _l, _m, _z) in movs:
        c.append('INSERT INTO dbo.Vac_Movimiento (IdColaborador, Tipo, Estado, FechaDesde, FechaHasta, Dias, Horas, '
                 'Minutos, BeneficioDos, Observacion, Origen, Orden, Usuario_Registro) SELECT IdColaborador, '
                 "%s, %s, %s, %s, %s, %s, %s, %d, %s, 'EXCEL', %d, 'CARGA_EXCEL' FROM dbo.Vac_Colaborador WHERE FilaExcel = %d;"
                 % (txt(tipo), txt(est), fecha(d), fecha(h), num(dias), num(horas), num(mins), ben, txt(obs), fila, n))
    for (d, desc) in feriados:
        c.append("INSERT INTO dbo.Vac_FeriadoHistorico (Fecha, Descripcion, Usuario_Registro) "
                 "SELECT %s, %s, 'CARGA_EXCEL' WHERE NOT EXISTS (SELECT 1 FROM dbo.Vac_FeriadoHistorico WHERE Fecha = %s);"
                 % (fecha(d), txt(desc), fecha(d)))
    c += ['COMMIT;',
          'DECLARE @C INT = (SELECT COUNT(*) FROM dbo.Vac_Colaborador), '
          "@M INT = (SELECT COUNT(*) FROM dbo.Vac_Movimiento WHERE Origen = 'EXCEL'), "
          '@H INT = (SELECT COUNT(*) FROM dbo.Vac_FeriadoHistorico);',
          "PRINT 'Carga: ' + CONVERT(VARCHAR(10), @C) + ' colaboradores, ' + CONVERT(VARCHAR(10), @M) + ' movimientos, '"
          " + CONVERT(VARCHAR(10), @H) + ' feriados historicos.';"]
    with open(os.path.join(salida, 'vacaciones-carga.sql'), 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(c) + '\n')

    # -------------------------------------------------------- sueldos.sql
    sueldos = []
    for i, r in enumerate(libro['COLABORADORES'].iter_rows(min_row=FILA_INICIO, values_only=True)):
        if r[2] and r[9] not in (None, '') and float(r[9]) > 0:
            sueldos.append((str(r[1]).strip(), r[9]))
    u = ["/* Generado por docs/sql/generar-carga-vacaciones.py el %s. NO VERSIONAR: sueldos. */"
         % datetime.date.today().isoformat(),
         'SET NOCOUNT ON;', 'SET XACT_ABORT ON;', 'SET QUOTED_IDENTIFIER ON;',
         'BEGIN TRANSACTION;',
         "DELETE FROM dbo.Vac_Sueldo WHERE Origen = 'EXCEL';"]
    for (ced, monto) in sueldos:
        u.append("INSERT INTO dbo.Vac_Sueldo (Cedula, Monto, FechaVigenciaDesde, Origen, Usuario_Registro) "
                 "VALUES (%s, %s, '19000101', 'EXCEL', 'CARGA_EXCEL');" % (txt(ced), num(monto)))
    u += ['COMMIT;',
          "DECLARE @S INT = (SELECT COUNT(*) FROM dbo.Vac_Sueldo WHERE Origen = 'EXCEL');",
          "PRINT 'Sueldos: ' + CONVERT(VARCHAR(10), @S) + ' cargados desde el Excel.';"]
    with open(os.path.join(salida, 'vacaciones-sueldos.sql'), 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(u) + '\n')

    # ------------------------------------------------------- prueba-oro.sql
    p = ["/* Generado por docs/sql/generar-carga-vacaciones.py. NO VERSIONAR. Prueba de oro del motor. */",
         'SET NOCOUNT ON;', 'SET QUOTED_IDENTIFIER ON;',
         "DECLARE @F DATE = %s;" % fecha(f_rep),
         'CREATE TABLE #EspSaldo (FilaExcel INT PRIMARY KEY, Disponible FLOAT, Usado FLOAT, Proporcional FLOAT, '
         'Saldo FLOAT, PorVencer FLOAT, Vencidos FLOAT, Gozado FLOAT, Estado VARCHAR(12));',
         'CREATE TABLE #EspPer (FilaExcel INT, N INT, Dias FLOAT, Tomado FLOAT, Estado VARCHAR(20), Saldo FLOAT, '
         'Vencido FLOAT, PRIMARY KEY (FilaExcel, N));',
         'CREATE TABLE #EspMov (Orden INT PRIMARY KEY, Descuenta FLOAT, Situacion VARCHAR(20), Cruces INT);']
    for r in libro['SALDOS'].iter_rows(min_row=FILA_INICIO, values_only=True):
        if not r[1]:
            continue
        p.append('INSERT INTO #EspSaldo VALUES (%d, %s, %s, %s, %s, %s, %s, %s, %s);'
                 % (r[0], num(r[13]), num(r[14]), num(r[15]), num(r[16]), num(r[17]), num(r[19]), num(r[20]),
                    "'" + ESTADOS_SALDO[r[21]] + "'"))
    for r in libro['CALCULO_PERIODOS'].iter_rows(min_row=FILA_INICIO, values_only=True):
        if r[2] is not True:
            continue
        estado = r[13]
        p.append("INSERT INTO #EspPer VALUES (%d, %d, %s, %s, '%s', %s, %s);"
                 % (r[0], r[1], num(r[7]), num(r[11]), estado, num(r[14]), num(r[15])))
    for (fila, n, tipo, est, d, h, dias, horas, mins, ben, obs, l, m, z) in movs:
        sit = 'DESCUENTA' if str(m).startswith('DESCUENTA') else SITUACIONES[m]
        p.append("INSERT INTO #EspMov VALUES (%d, %s, '%s', %d);" % (fila, num(l or 0), sit, z or 0))

    p.append(r"""
CREATE TABLE #ResSaldo (IdColaborador INT, EstadoColaborador VARCHAR(10), FechaCalculo DATE, Cumplidos INT, DiasPorAnio DECIMAL(12,4),
    InicioPeriodo DATE, ProximoAniversario DATE, AjusteCorte DECIMAL(12,4), Disponible DECIMAL(12,4), UsadoEnCurso DECIMAL(12,4),
    Proporcional DECIMAL(12,4), Saldo DECIMAL(12,2), PorVencer DECIMAL(12,4), ProximoVencimiento DATE, Vencidos DECIMAL(12,4),
    TotalGozado DECIMAL(12,4), Estado VARCHAR(12));
CREATE TABLE #ResPer (IdColaborador INT, N INT, Periodo VARCHAR(9), Inicio DATE, Fin DATE, DiasGanados DECIMAL(12,4),
    DiasTomados DECIMAL(12,4), DiasVencidos DECIMAL(12,4), Saldo DECIMAL(12,4), Vence DATE, Estado VARCHAR(20));
CREATE TABLE #ResMov (IdMovimiento INT, IdColaborador INT, Tipo VARCHAR(25), Estado VARCHAR(15), Desde DATE, Hasta DATE,
    Situacion VARCHAR(20), AnioServicio INT, DiasEquivalentes DECIMAL(12,4), DiasBase DECIMAL(12,4),
    DiasFinDeSemana DECIMAL(12,4), DiasDescuenta DECIMAL(12,4), Cruces INT);
INSERT INTO #ResSaldo EXEC dbo.Sp_RTA_VacCalcular @FechaCalculo = @F, @Salida = 'SALDOS';
INSERT INTO #ResPer   EXEC dbo.Sp_RTA_VacCalcular @FechaCalculo = @F, @Salida = 'PERIODOS';
INSERT INTO #ResMov   EXEC dbo.Sp_RTA_VacCalcular @FechaCalculo = @F, @Salida = 'MOVIMIENTOS';

DECLARE @Fallos INT = 0, @Rc INT, @NMov INT = (SELECT COUNT(*) FROM #EspMov),
        @NPer INT = (SELECT COUNT(*) FROM #EspPer), @NSaldo INT = (SELECT COUNT(*) FROM #EspSaldo);

/* 1. Movimientos: lo que descuenta cada solicitud y su situacion. */
SELECT e.Orden, e.Descuenta AS Excel, m.DiasDescuenta AS Motor, e.Situacion AS SitExcel, m.Situacion AS SitMotor,
       e.Cruces AS CrucesExcel, m.Cruces AS CrucesMotor
INTO #DifMov
FROM #EspMov e
LEFT JOIN dbo.Vac_Movimiento v ON v.Origen = 'EXCEL' AND v.Orden = e.Orden
LEFT JOIN #ResMov m ON m.IdMovimiento = v.IdMovimiento
WHERE m.IdMovimiento IS NULL OR ABS(e.Descuenta - m.DiasDescuenta) > 0.0001 OR e.Situacion <> m.Situacion OR e.Cruces <> m.Cruces;
SET @Rc = (SELECT COUNT(*) FROM #DifMov);
IF @Rc = 0 PRINT 'PASS 1: las ' + CONVERT(VARCHAR(10), @NMov) + ' solicitudes descuentan lo mismo que en el Excel.'
ELSE BEGIN PRINT 'FAIL 1: ' + CONVERT(VARCHAR(10), @Rc) + ' solicitudes difieren (fila de REGISTRO):'; SET @Fallos += 1;
    SELECT TOP 30 * FROM #DifMov ORDER BY Orden; END

/* 2. Periodos: dias, tomado, estado, saldo y vencido de cada periodo abierto. */
SELECT e.FilaExcel, e.N, e.Tomado AS TomadoExcel, p.DiasTomados AS TomadoMotor, e.Estado AS EstExcel, p.Estado AS EstMotor,
       e.Saldo AS SaldoExcel, p.Saldo AS SaldoMotor
INTO #DifPer
FROM #EspPer e
LEFT JOIN dbo.Vac_Colaborador c ON c.FilaExcel = e.FilaExcel
LEFT JOIN #ResPer p ON p.IdColaborador = c.IdColaborador AND p.N = e.N
WHERE p.N IS NULL OR ABS(e.Dias - p.DiasGanados) > 0.0001 OR ABS(e.Tomado - p.DiasTomados) > 0.0001
   OR e.Estado <> p.Estado OR ABS(e.Saldo - p.Saldo) > 0.0001 OR ABS(e.Vencido - p.DiasVencidos) > 0.0001;
SET @Rc = (SELECT COUNT(*) FROM #DifPer)
        + (SELECT COUNT(*) FROM #ResPer p JOIN dbo.Vac_Colaborador c ON c.IdColaborador = p.IdColaborador
           WHERE NOT EXISTS (SELECT 1 FROM #EspPer e WHERE e.FilaExcel = c.FilaExcel AND e.N = p.N));
IF @Rc = 0 PRINT 'PASS 2: los ' + CONVERT(VARCHAR(10), @NPer) + ' periodos cuadran con CALCULO_PERIODOS.'
ELSE BEGIN PRINT 'FAIL 2: ' + CONVERT(VARCHAR(10), @Rc) + ' periodos difieren (N de COLABORADORES, periodo):'; SET @Fallos += 1;
    SELECT TOP 30 * FROM #DifPer ORDER BY FilaExcel, N; END

/* 3. Saldos por persona, al centesimo. */
SELECT e.FilaExcel, e.Saldo AS SaldoExcel, s.Saldo AS SaldoMotor, e.Disponible AS DispExcel, s.Disponible AS DispMotor,
       e.Proporcional AS PropExcel, s.Proporcional AS PropMotor, e.Estado AS EstExcel, s.Estado AS EstMotor
INTO #DifSaldo
FROM #EspSaldo e
LEFT JOIN dbo.Vac_Colaborador c ON c.FilaExcel = e.FilaExcel
LEFT JOIN #ResSaldo s ON s.IdColaborador = c.IdColaborador
WHERE s.IdColaborador IS NULL OR ABS(e.Saldo - s.Saldo) > 0.005 OR ABS(e.Disponible - s.Disponible) > 0.0001
   OR ABS(e.Usado - s.UsadoEnCurso) > 0.0001 OR ABS(e.Proporcional - s.Proporcional) > 0.0001
   OR ABS(e.PorVencer - s.PorVencer) > 0.0001 OR ABS(e.Vencidos - s.Vencidos) > 0.0001
   OR ABS(e.Gozado - s.TotalGozado) > 0.0001 OR e.Estado <> s.Estado;
SET @Rc = (SELECT COUNT(*) FROM #DifSaldo);
IF @Rc = 0 PRINT 'PASS 3: los ' + CONVERT(VARCHAR(10), @NSaldo) + ' saldos cuadran con SALDOS.'
ELSE BEGIN PRINT 'FAIL 3: ' + CONVERT(VARCHAR(10), @Rc) + ' personas difieren (N de COLABORADORES):'; SET @Fallos += 1;
    SELECT TOP 30 * FROM #DifSaldo ORDER BY FilaExcel; END

/* 4. Totales del DASHBOARD (solo activos). */
SELECT 'Excel' AS Lado, SUM(Disponible) AS Disponible, SUM(Saldo) AS Saldo FROM #EspSaldo e
    JOIN dbo.Vac_Colaborador c ON c.FilaExcel = e.FilaExcel WHERE c.Estado = 'ACTIVO'
UNION ALL
SELECT 'Motor', SUM(Disponible), SUM(Saldo) FROM #ResSaldo WHERE EstadoColaborador = 'ACTIVO';

PRINT CASE WHEN @Fallos = 0 THEN 'VEREDICTO: OK' ELSE 'VEREDICTO: FALLA (' + CONVERT(VARCHAR(10), @Fallos) + ')' END;
""")
    with open(os.path.join(salida, 'vacaciones-prueba-oro.sql'), 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(p) + '\n')

    print('Colaboradores: %d  Movimientos: %d  Feriados historicos: %d  Fecha de calculo: %s'
          % (len(colabs), len(movs), len(feriados), f_rep.date().isoformat()))


if __name__ == '__main__':
    main()
